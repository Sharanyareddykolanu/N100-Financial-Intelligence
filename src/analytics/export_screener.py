import os
import sys
import sqlite3

PROJECT_ROOT = os.path.abspath(
    os.path.join(os.path.dirname(__file__), "..", "..")
)

if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from src.analytics.composite_score import calculate_composite_score
from src.screener.presets import (
    quality_compounder,
    value_pick,
    growth_accelerator,
    dividend_champion,
    debt_free_blue_chip,
    turnaround_watch,
)


DB_PATH = os.path.join(PROJECT_ROOT, "nifty100.db")
OUTPUT_DIR = os.path.join(PROJECT_ROOT, "output")
OUTPUT_PATH = os.path.join(
    OUTPUT_DIR,
    "screener_output.xlsx",
)


def load_data():
    conn = sqlite3.connect(DB_PATH)

    query = """
    SELECT
        fr.company_id,
        fr.year,
        fr.return_on_equity_pct,
        fr.net_profit_margin_pct,
        fr.debt_to_equity,
        fr.interest_coverage,
        fr.free_cash_flow_cr,
        fr.cash_from_operations_cr,
        fr.revenue_cagr_5yr,
        fr.pat_cagr_5yr,

        pl.sales,
        pl.operating_profit,
        pl.net_profit,

        bs.total_equity,
        bs.debt,

        cv.pe_ratio,
        cv.pb_ratio,
        cv.dividend_yield,

        c.company_name,
        c.ticker,

        cs.sector

    FROM financial_ratios fr

    LEFT JOIN company_profit_loss pl
        ON fr.company_id = pl.company_id
        AND fr.year = pl.year

    LEFT JOIN company_balance_sheet bs
        ON fr.company_id = bs.company_id
        AND fr.year = bs.year

    LEFT JOIN company_valuation cv
        ON fr.company_id = cv.company_id
        AND fr.year = cv.year

    LEFT JOIN companies c
        ON fr.company_id = c.company_id

    LEFT JOIN company_sector cs
        ON fr.company_id = cs.company_id

    ORDER BY fr.company_id, fr.year
    """

    df = pd.read_sql_query(query, conn)

    conn.close()

    return df


def add_payout_ratio(df):
    conn = sqlite3.connect(DB_PATH)

    query = """
    SELECT
        company_id,
        year,
        dividend_payout_ratio_pct
    FROM financial_ratios
    """

    payout = pd.read_sql_query(query, conn)

    conn.close()

    payout = payout.drop_duplicates(
        subset=["company_id", "year"]
    )

    df = df.merge(
        payout,
        on=["company_id", "year"],
        how="left",
        suffixes=("", "_payout"),
    )

    return df


def latest_year(df):
    if df.empty:
        return df.copy()

    return (
        df.sort_values(
            ["company_id", "year"]
        )
        .groupby(
            "company_id",
            as_index=False
        )
        .tail(1)
        .copy()
    )


def add_turnaround_metrics(df):
    data = df.sort_values(
        ["company_id", "year"]
    ).copy()

    previous_sales = (
        data.groupby("company_id")["sales"]
        .shift(3)
    )

    previous_de = (
        data.groupby("company_id")[
            "debt_to_equity"
        ].shift(1)
    )

    data["revenue_cagr_3yr"] = (
        (
            data["sales"] / previous_sales
        ) ** (1 / 3)
        - 1
    ) * 100

    data.loc[
        (data["sales"] <= 0)
        | (previous_sales <= 0),
        "revenue_cagr_3yr"
    ] = pd.NA

    data["previous_de"] = previous_de

    return data


def prepare_data():
    print("Loading database data...")

    df = load_data()

    print(f"Loaded rows: {len(df)}")
    print(
        f"Companies: {df['company_id'].nunique()}"
    )

    df = add_payout_ratio(df)

    df = add_turnaround_metrics(df)

    print("Calculating composite scores...")

    df = calculate_composite_score(df)

    print(
        "Composite scores available: "
        f"{df['composite_score'].notna().sum()}"
    )

    return df


def get_presets(df):
    return {
        "Quality Compounder": quality_compounder(df),
        "Value Pick": value_pick(df),
        "Growth Accelerator": growth_accelerator(df),
        "Dividend Champion": dividend_champion(df),
        "Debt-Free Blue Chip": debt_free_blue_chip(df),
        "Turnaround Watch": turnaround_watch(df),
    }


def calculate_roce(df):
    denominator = (
        df["total_equity"]
        + df["debt"]
    )

    return (
        df["operating_profit"]
        / denominator
        * 100
    ).where(
        denominator > 0
    )


def calculate_fcf_cagr(df):
    data = df.sort_values(
        ["company_id", "year"]
    ).copy()

    previous_fcf = (
        data.groupby("company_id")[
            "free_cash_flow_cr"
        ].shift(5)
    )

    current_fcf = data[
        "free_cash_flow_cr"
    ]

    data["fcf_cagr_5yr"] = (
        (
            current_fcf
            / previous_fcf
        ) ** (1 / 5)
        - 1
    ) * 100

    data.loc[
        (current_fcf <= 0)
        | (previous_fcf <= 0),
        "fcf_cagr_5yr"
    ] = pd.NA

    return data


def prepare_export_columns(df):
    result = df.copy()

    result["roce_pct"] = calculate_roce(
        result
    )

    result = calculate_fcf_cagr(
        result
    )

    result["cfo_pat_ratio"] = (
        result["cash_from_operations_cr"]
        / result["net_profit"]
        * 100
    ).where(
        result["net_profit"] != 0
    )

    result["fcf_positive_flag"] = (
        result["free_cash_flow_cr"] > 0
    ).astype(int)

    result = result.sort_values(
        "composite_score",
        ascending=False
    )

    columns = [
        "company_id",
        "company_name",
        "ticker",
        "sector",
        "year",
        "return_on_equity_pct",
        "roce_pct",
        "net_profit_margin_pct",
        "free_cash_flow_cr",
        "fcf_cagr_5yr",
        "cash_from_operations_cr",
        "net_profit",
        "cfo_pat_ratio",
        "fcf_positive_flag",
        "revenue_cagr_5yr",
        "pat_cagr_5yr",
        "debt_to_equity",
        "interest_coverage",
        "composite_score",
        "sales",
    ]

    result = result[
        columns
    ].copy()

    return result.sort_values(
        "composite_score",
        ascending=False
    )


def write_excel(presets):
    os.makedirs(
        OUTPUT_DIR,
        exist_ok=True
    )

    with pd.ExcelWriter(
        OUTPUT_PATH,
        engine="openpyxl"
    ) as writer:

        for preset_name, preset_df in presets.items():

            latest = latest_year(
                preset_df
            )

            export_df = prepare_export_columns(
                latest
            )

            export_df.to_excel(
                writer,
                sheet_name=preset_name[:31],
                index=False
            )

            print(
                f"{preset_name}: "
                f"{len(export_df)} companies exported"
            )

    return OUTPUT_PATH


def apply_formatting(path):
    workbook = load_workbook(path)

    green_fill = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE"
    )

    red_fill = PatternFill(
        fill_type="solid",
        fgColor="FFC7CE"
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7"
    )

    rules = {
        "Quality Compounder": {
            "return_on_equity_pct": ("gt", 15),
            "debt_to_equity": ("lt", 1.0),
            "free_cash_flow_cr": ("gt", 0),
            "revenue_cagr_5yr": ("gt", 10),
        },

        "Value Pick": {
            "debt_to_equity": ("lt", 2.0),
        },

        "Growth Accelerator": {
            "pat_cagr_5yr": ("gt", 20),
            "revenue_cagr_5yr": ("gt", 15),
            "debt_to_equity": ("lt", 2.0),
        },

        "Dividend Champion": {
            "free_cash_flow_cr": ("gt", 0),
        },

        "Debt-Free Blue Chip": {
            "debt_to_equity": ("eq", 0),
            "return_on_equity_pct": ("gt", 12),
            "sales": ("gt", 5000),
        },

        "Turnaround Watch": {
            "revenue_cagr_3yr": ("gt", 10),
            "free_cash_flow_cr": ("gt", 0),
        },
    }

    for sheet_name in workbook.sheetnames:

        ws = workbook[sheet_name]

        for cell in ws[1]:
            cell.fill = header_fill

        headers = {
            cell.value: cell.column
            for cell in ws[1]
        }

        preset_rules = rules.get(
            sheet_name,
            {}
        )

        for column_name, (
            operator,
            threshold
        ) in preset_rules.items():

            if column_name not in headers:
                continue

            column_number = headers[
                column_name
            ]

            for row_number in range(
                2,
                ws.max_row + 1
            ):

                cell = ws.cell(
                    row=row_number,
                    column=column_number
                )

                if cell.value is None:
                    continue

                try:
                    value = float(
                        cell.value
                    )
                except (
                    TypeError,
                    ValueError
                ):
                    cell.fill = red_fill
                    continue

                if operator == "gt":
                    passed = (
                        value > threshold
                    )

                elif operator == "lt":
                    passed = (
                        value < threshold
                    )

                elif operator == "eq":
                    passed = (
                        value == threshold
                    )

                else:
                    passed = False

                cell.fill = (
                    green_fill
                    if passed
                    else red_fill
                )

        for row in ws.iter_rows(
            min_row=2
        ):
            for cell in row:
                if isinstance(
                    cell.value,
                    (int, float)
                ):
                    cell.number_format = "0.00"

        ws.freeze_panes = "A2"

        for column_cells in ws.columns:

            max_length = 0

            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:

                if cell.value is None:
                    continue

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

            ws.column_dimensions[
                column_letter
            ].width = min(
                max(max_length + 2, 10),
                35
            )

    workbook.save(path)


def print_summary(presets):

    print()
    print("=" * 70)
    print("DAY 17 EXPORT SUMMARY")
    print("=" * 70)

    for name, df in presets.items():

        latest = latest_year(df)

        print(
            f"{name}: "
            f"{len(latest)} companies"
        )

    print()
    print(
        f"Output file: {OUTPUT_PATH}"
    )

    print(
        "Excel export completed successfully."
    )


def main():

    print("=" * 70)
    print("DAY 17 - COMPOSITE SCORE & EXCEL EXPORT")
    print("=" * 70)

    try:

        df = prepare_data()

        presets = get_presets(df)

        output_path = write_excel(
            presets
        )

        apply_formatting(
            output_path
        )

        print_summary(
            presets
        )

    except Exception as exc:

        print()
        print("=" * 70)
        print("EXPORT FAILED")
        print("=" * 70)

        print(
            f"{type(exc).__name__}: {exc}"
        )

        raise


if __name__ == "__main__":
    main()
