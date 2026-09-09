from pathlib import Path
import sqlite3

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# PATHS
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]

DB_PATH = PROJECT_ROOT / "nifty100.db"
PEER_FILE = PROJECT_ROOT / "peer_groups.xlsx"
OUTPUT_FILE = PROJECT_ROOT / "output" / "peer_comparison.xlsx"


# ============================================================
# METRICS
# ============================================================

METRICS = {
    "ROE": "ROE",
    "ROCE": "ROCE",
    "NPM": "Net Profit Margin",
    "D/E": "D/E",
    "FCF": "FCF",
    "PAT CAGR 5yr": "PAT CAGR 5yr",
    "Revenue CAGR 5yr": "Revenue CAGR 5yr",
    "EPS CAGR 5yr": "EPS CAGR 5yr",
    "Interest Coverage": "Interest Coverage",
    "Asset Turnover": "Asset Turnover",
}


# ============================================================
# LOAD PEER GROUPS
# ============================================================

def load_peer_groups():
    print("Loading peer_groups.xlsx...")

    df = pd.read_excel(PEER_FILE)

    df.columns = [
        str(column).strip().lower().replace(" ", "_")
        for column in df.columns
    ]

    required = {
        "id",
        "peer_group_name",
        "company_id",
        "is_benchmark",
    }

    missing = required - set(df.columns)

    if missing:
        raise ValueError(
            f"Missing peer group columns: {missing}"
        )

    df["company_id"] = (
        df["company_id"]
        .astype(str)
        .str.strip()
        .str.upper()
    )

    df["peer_group_name"] = (
        df["peer_group_name"]
        .astype(str)
        .str.strip()
    )

    return df


# ============================================================
# LOAD COMPANY NAMES
# ============================================================

def load_company_names(conn):

    tables = pd.read_sql_query(
        """
        SELECT name
        FROM sqlite_master
        WHERE type = 'table'
        ORDER BY name
        """,
        conn,
    )["name"].tolist()

    possible_tables = [
        "companies",
        "company",
        "company_master",
        "company_details",
    ]

    for table in possible_tables:

        if table not in tables:
            continue

        columns = pd.read_sql_query(
            f'PRAGMA table_info("{table}")',
            conn,
        )["name"].tolist()

        id_column = None
        name_column = None

        for column in columns:

            lower = column.lower()

            if lower in {
                "company_id",
                "id",
                "ticker",
                "symbol",
            }:
                id_column = column

            if lower in {
                "company_name",
                "name",
                "company",
            }:
                name_column = column

        if id_column and name_column:

            print(
                f"Company names found in table: {table}"
            )

            query = f"""
                SELECT
                    "{id_column}" AS company_id,
                    "{name_column}" AS company_name
                FROM "{table}"
            """

            result = pd.read_sql_query(
                query,
                conn,
            )

            result["company_id"] = (
                result["company_id"]
                .astype(str)
                .str.strip()
                .str.upper()
            )

            result["company_name"] = (
                result["company_name"]
                .astype(str)
                .str.strip()
            )

            return result

    print(
        "Company names table not found."
    )

    return pd.DataFrame(
        columns=[
            "company_id",
            "company_name",
        ]
    )


# ============================================================
# LOAD PEER DATA
# ============================================================

def load_peer_data():

    conn = sqlite3.connect(DB_PATH)

    peer_data = pd.read_sql_query(
        """
        SELECT
            company_id,
            peer_group_name,
            metric,
            value,
            percentile_rank,
            year
        FROM peer_percentiles
        """,
        conn,
    )

    company_names = load_company_names(conn)

    conn.close()

    peer_data["company_id"] = (
        peer_data["company_id"]
        .astype(str)
        .str.strip()
        .str.upper()
    )

    peer_data["peer_group_name"] = (
        peer_data["peer_group_name"]
        .astype(str)
        .str.strip()
    )

    return peer_data, company_names


# ============================================================
# GET LATEST DATA
# ============================================================

def get_latest_peer_data(peer_data):

    data = peer_data.copy()

    data = data.sort_values(
        [
            "company_id",
            "metric",
            "year",
        ]
    )

    latest = (
        data
        .groupby(
            [
                "company_id",
                "peer_group_name",
                "metric",
            ],
            as_index=False,
        )
        .tail(1)
    )

    return latest


# ============================================================
# BUILD GROUP DATAFRAME
# ============================================================

def build_group_dataframe(
    group_name,
    peer_data,
    company_names,
):

    data = peer_data[
        peer_data["peer_group_name"]
        == group_name
    ].copy()

    data = get_latest_peer_data(data)

    # --------------------------------------------------------
    # VALUES
    # --------------------------------------------------------

    values = data.pivot_table(
        index="company_id",
        columns="metric",
        values="value",
        aggfunc="last",
    )

    # --------------------------------------------------------
    # PERCENTILES
    # --------------------------------------------------------

    percentiles = data.pivot_table(
        index="company_id",
        columns="metric",
        values="percentile_rank",
        aggfunc="last",
    )

    # --------------------------------------------------------
    # RENAME VALUE COLUMNS
    # --------------------------------------------------------

    value_rename = {}

    for display_name, source_name in METRICS.items():

        value_rename[source_name] = (
            f"{display_name} Value"
        )

    values = values.rename(
        columns=value_rename
    )

    # --------------------------------------------------------
    # RENAME PERCENTILE COLUMNS
    # --------------------------------------------------------

    percentile_rename = {}

    for display_name, source_name in METRICS.items():

        percentile_rename[source_name] = (
            f"{display_name} Percentile"
        )

    percentiles = percentiles.rename(
        columns=percentile_rename
    )

    # --------------------------------------------------------
    # COMBINE
    # --------------------------------------------------------

    result = values.join(
        percentiles,
        how="outer",
    )

    result = result.reset_index()

    result["company_id"] = (
        result["company_id"]
        .astype(str)
        .str.strip()
        .str.upper()
    )

    # --------------------------------------------------------
    # COMPANY NAME
    # --------------------------------------------------------

    if not company_names.empty:

        names = company_names.copy()

        names["company_id"] = (
            names["company_id"]
            .astype(str)
            .str.strip()
            .str.upper()
        )

        names = names.drop_duplicates(
            subset=["company_id"]
        )

        result = result.merge(
            names[
                [
                    "company_id",
                    "company_name",
                ]
            ],
            on="company_id",
            how="left",
        )

    else:

        result["company_name"] = (
            result["company_id"]
        )

    result["company_name"] = (
        result["company_name"]
        .fillna(result["company_id"])
    )

    # --------------------------------------------------------
    # COLUMN ORDER
    # --------------------------------------------------------

    ordered_columns = [
        "company_id",
        "company_name",
    ]

    for metric in METRICS:

        ordered_columns.append(
            f"{metric} Value"
        )

    for metric in METRICS:

        ordered_columns.append(
            f"{metric} Percentile"
        )

    # Add missing columns.
    for column in ordered_columns:

        if column not in result.columns:

            result[column] = pd.NA

    result = result[
        ordered_columns
    ]

    return result


# ============================================================
# EXPORT EXCEL
# ============================================================

def export_excel(
    peer_data,
    company_names,
    peer_group_table,
):

    OUTPUT_FILE.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    # IMPORTANT:
    # Get actual peer group names from the DataFrame.
    peer_groups = (
        peer_group_table[
            "peer_group_name"
        ]
        .dropna()
        .astype(str)
        .str.strip()
        .unique()
        .tolist()
    )

    print(
        f"Peer groups to export: "
        f"{len(peer_groups)}"
    )

    # --------------------------------------------------------
    # First create workbook
    # --------------------------------------------------------

    with pd.ExcelWriter(
        OUTPUT_FILE,
        engine="openpyxl",
    ) as writer:

        for group_name in peer_groups:

            print(
                f"Exporting: {group_name}"
            )

            result = build_group_dataframe(
                group_name,
                peer_data,
                company_names,
            )

            sheet_name = group_name[:31]

            result.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
            )

    # --------------------------------------------------------
    # Re-open for formatting
    # --------------------------------------------------------

    workbook = load_workbook(
        OUTPUT_FILE
    )

    # --------------------------------------------------------
    # Fills
    # --------------------------------------------------------

    green_fill = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE",
    )

    yellow_fill = PatternFill(
        fill_type="solid",
        fgColor="FFEB9C",
    )

    red_fill = PatternFill(
        fill_type="solid",
        fgColor="FFC7CE",
    )

    gold_fill = PatternFill(
        fill_type="solid",
        fgColor="FFD966",
    )

    median_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAD3",
    )

    # --------------------------------------------------------
    # Format every peer group
    # --------------------------------------------------------

    for group_name in peer_groups:

        sheet_name = group_name[:31]

        ws = workbook[sheet_name]

        # ----------------------------------------------------
        # Header
        # ----------------------------------------------------

        for cell in ws[1]:

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        ws.row_dimensions[1].height = 35

        # ----------------------------------------------------
        # Benchmark company
        # ----------------------------------------------------

        benchmark_rows = peer_group_table[
            peer_group_table[
                "peer_group_name"
            ]
            == group_name
        ].copy()

        benchmark_rows = benchmark_rows[
            benchmark_rows[
                "is_benchmark"
            ] == True
        ]

        benchmark_ids = set(
            benchmark_rows[
                "company_id"
            ]
            .astype(str)
            .str.strip()
            .str.upper()
        )

        # ----------------------------------------------------
        # Percentile columns
        # ----------------------------------------------------

        percentile_columns = []

        for column_index in range(
            1,
            ws.max_column + 1,
        ):

            header = ws.cell(
                row=1,
                column=column_index,
            ).value

            if (
                header is not None
                and str(header).endswith(
                    "Percentile"
                )
            ):

                percentile_columns.append(
                    column_index
                )

        # ----------------------------------------------------
        # Data rows
        # ----------------------------------------------------

        data_end = ws.max_row

        for row_index in range(
            2,
            data_end + 1,
        ):

            company_id = ws.cell(
                row=row_index,
                column=1,
            ).value

            if company_id is None:
                continue

            company_id = (
                str(company_id)
                .strip()
                .upper()
            )

            # ------------------------------------------------
            # Benchmark row
            # ------------------------------------------------

            if company_id in benchmark_ids:

                for column_index in range(
                    1,
                    ws.max_column + 1,
                ):

                    cell = ws.cell(
                        row=row_index,
                        column=column_index,
                    )

                    cell.fill = gold_fill

                    cell.font = Font(
                        bold=True
                    )

            # ------------------------------------------------
            # Percentile colours
            # ------------------------------------------------

            for column_index in percentile_columns:

                cell = ws.cell(
                    row=row_index,
                    column=column_index,
                )

                if (
                    cell.value is None
                    or cell.value == ""
                ):
                    continue

                try:

                    percentile = float(
                        cell.value
                    )

                except (
                    ValueError,
                    TypeError,
                ):

                    continue

                if percentile >= 0.75:

                    cell.fill = green_fill

                elif percentile <= 0.25:

                    cell.fill = red_fill

                else:

                    cell.fill = yellow_fill

                cell.number_format = "0.0%"

        # ----------------------------------------------------
        # MEDIAN ROW
        # ----------------------------------------------------

        median_row = ws.max_row + 1

        ws.cell(
            row=median_row,
            column=1,
            value="PEER GROUP MEDIAN",
        )

        ws.cell(
            row=median_row,
            column=2,
            value=group_name,
        )

        # ----------------------------------------------------
        # Median for metric VALUE columns
        # ----------------------------------------------------

        for column_index in range(
            3,
            ws.max_column + 1,
        ):

            header = ws.cell(
                row=1,
                column=column_index,
            ).value

            if (
                header is None
                or not str(header).endswith(
                    "Value"
                )
            ):

                continue

            numeric_values = []

            for row_index in range(
                2,
                median_row,
            ):

                value = ws.cell(
                    row=row_index,
                    column=column_index,
                ).value

                if (
                    value is None
                    or value == ""
                ):

                    continue

                try:

                    numeric_values.append(
                        float(value)
                    )

                except (
                    ValueError,
                    TypeError,
                ):

                    continue

            if numeric_values:

                median_value = float(
                    pd.Series(
                        numeric_values
                    ).median()
                )

                ws.cell(
                    row=median_row,
                    column=column_index,
                    value=median_value,
                )

        # ----------------------------------------------------
        # Median row formatting
        # ----------------------------------------------------

        for column_index in range(
            1,
            ws.max_column + 1,
        ):

            cell = ws.cell(
                row=median_row,
                column=column_index,
            )

            cell.fill = median_fill

            cell.font = Font(
                bold=True
            )

        # ----------------------------------------------------
        # Freeze panes
        # ----------------------------------------------------

        ws.freeze_panes = "C2"

        # ----------------------------------------------------
        # Auto filter
        # ----------------------------------------------------

        ws.auto_filter.ref = (
            f"A1:"
            f"{get_column_letter(ws.max_column)}"
            f"{median_row - 1}"
        )

        # ----------------------------------------------------
        # Column widths
        # ----------------------------------------------------

        for column_index in range(
            1,
            ws.max_column + 1,
        ):

            letter = get_column_letter(
                column_index
            )

            if column_index == 1:

                width = 18

            elif column_index == 2:

                width = 25

            else:

                width = 20

            ws.column_dimensions[
                letter
            ].width = width

    # --------------------------------------------------------
    # Save workbook
    # --------------------------------------------------------

    workbook.save(
        OUTPUT_FILE
    )

    print()
    print(
        f"Output file: {OUTPUT_FILE}"
    )


# ============================================================
# VALIDATION
# ============================================================

def validate_excel(peer_group_table):

    print()
    print("=" * 70)
    print("DAY 20 VALIDATION")
    print("=" * 70)

    workbook = load_workbook(
        OUTPUT_FILE,
        data_only=False,
    )

    # --------------------------------------------------------
    # Expected peer groups
    # --------------------------------------------------------

    expected_groups = (
        peer_group_table[
            "peer_group_name"
        ]
        .dropna()
        .astype(str)
        .str.strip()
        .unique()
        .tolist()
    )

    expected_groups = sorted(
        expected_groups
    )

    actual_groups = sorted(
        workbook.sheetnames
    )

    print(
        f"Sheets generated: "
        f"{len(actual_groups)}"
    )

    print(
        f"Expected sheets: "
        f"{len(expected_groups)}"
    )

    if (
        len(actual_groups)
        == len(expected_groups)
        and actual_groups
        == [
            group[:31]
            for group in expected_groups
        ]
    ):

        print(
            "11 peer-group sheets: PASS"
        )

    else:

        print(
            "11 peer-group sheets: FAIL"
        )

        print(
            f"Actual: {actual_groups}"
        )

        print(
            f"Expected: {expected_groups}"
        )

    # --------------------------------------------------------
    # Expected columns
    # --------------------------------------------------------

    expected_columns = [
        "company_id",
        "company_name",
    ]

    for metric in METRICS:

        expected_columns.append(
            f"{metric} Value"
        )

    for metric in METRICS:

        expected_columns.append(
            f"{metric} Percentile"
        )

    print(
        f"Expected columns: "
        f"{len(expected_columns)}"
    )

    all_columns_ok = True

    for sheet_name in workbook.sheetnames:

        ws = workbook[sheet_name]

        actual_columns = [
            ws.cell(
                row=1,
                column=column_index,
            ).value
            for column_index in range(
                1,
                ws.max_column + 1,
            )
        ]

        if actual_columns != expected_columns:

            all_columns_ok = False

            print(
                f"Column check failed: "
                f"{sheet_name}"
            )

    if all_columns_ok:

        print(
            "20 metric columns + ID/name: PASS"
        )

    else:

        print(
            "Column structure: FAIL"
        )

    # --------------------------------------------------------
    # Benchmark check
    # --------------------------------------------------------

    benchmark_ok = True

    for group_name in expected_groups:

        sheet_name = group_name[:31]

        if sheet_name not in workbook.sheetnames:

            benchmark_ok = False

            print(
                f"Missing sheet: {group_name}"
            )

            continue

        ws = workbook[sheet_name]

        benchmark_rows = peer_group_table[
            peer_group_table[
                "peer_group_name"
            ]
            == group_name
        ]

        benchmark_rows = benchmark_rows[
            benchmark_rows[
                "is_benchmark"
            ] == True
        ]

        benchmark_ids = set(
            benchmark_rows[
                "company_id"
            ]
            .astype(str)
            .str.upper()
        )

        found = False

        for row_index in range(
            2,
            ws.max_row,
        ):

            company_id = str(
                ws.cell(
                    row=row_index,
                    column=1,
                ).value
            ).upper()

            if company_id in benchmark_ids:

                found = True
                break

        if not found:

            benchmark_ok = False

            print(
                f"Benchmark missing: "
                f"{group_name}"
            )

    if benchmark_ok:

        print(
            "Benchmark rows: PASS"
        )

    else:

        print(
            "Benchmark rows: FAIL"
        )

    workbook.close()

    print("=" * 70)


# ============================================================
# MAIN
# ============================================================

def main():

    print("=" * 70)
    print("DAY 20 - PEER COMPARISON EXCEL REPORT")
    print("=" * 70)

    # --------------------------------------------------------
    # Load peer group definition
    # --------------------------------------------------------

    peer_group_table = load_peer_groups()

    # --------------------------------------------------------
    # Load percentile data
    # --------------------------------------------------------

    peer_data, company_names = (
        load_peer_data()
    )

    print(
        f"Peer rows loaded: "
        f"{len(peer_data)}"
    )

    print(
        f"Peer groups: "
        f"{peer_data['peer_group_name'].nunique()}"
    )

    # --------------------------------------------------------
    # Export
    # --------------------------------------------------------

    export_excel(
        peer_data,
        company_names,
        peer_group_table,
    )

    # --------------------------------------------------------
    # Validate
    # --------------------------------------------------------

    validate_excel(
        peer_group_table
    )

    print()
    print("=" * 70)
    print("DAY 20 COMPLETED SUCCESSFULLY")
    print("=" * 70)


if __name__ == "__main__":
    main()